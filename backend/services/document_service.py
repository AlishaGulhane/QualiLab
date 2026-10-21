import base64
import io
from typing import Tuple
from fastapi import UploadFile, HTTPException
from PIL import Image



async def extract_content_paged(file: UploadFile) -> Tuple[str, list, list]:
    """
    Returns (content_type, pages, image_blocks)
    pages = [{"page_num": int, "text": str, "images": list}, ...]
    Used for paginated processing of PDFs.
    """
    filename = file.filename.lower()
    raw = await file.read()

    if filename.endswith(".pdf"):
        pages = extract_pdf_pages(raw)
        return "pdf", pages, []

    if filename.endswith(".docx"):
        text, _, images = _extract_docx(raw)
        return "docx", [{"page_num": 1, "text": text, "images": images}], []

    if filename.endswith((".jpg", ".jpeg", ".png", ".webp", ".gif", ".jfif")):
        _, _, images = _extract_image(raw, filename)
        return "image", [{"page_num": 1, "text": "", "images": images}], []

    if filename.endswith(".txt"):
        text = raw.decode("utf-8", errors="ignore")
        return "text", [{"page_num": 1, "text": text, "images": []}], []

    raise HTTPException(status_code=415, detail=f"Unsupported file type: {file.filename}")


def _tile_image_b64(b64_png: str, page: int, overlap: float = 0.1) -> list:
    """
    Split a base64 PNG into top and bottom tiles with overlap so the LLM is
    forced to focus on each half of the page rather than skimming the whole.
    Returns a list of image dicts (same schema as page renders).
    """
    img_bytes = base64.b64decode(b64_png)
    img = Image.open(io.BytesIO(img_bytes))
    w, h = img.size

    # Only tile if the image is tall enough to benefit (portrait pages)
    if h < w * 1.2:
        return [{"page": page, "data": b64_png, "mime": "image/png"}]

    overlap_px = int(h * overlap)
    mid = h // 2

    top = img.crop((0, 0, w, mid + overlap_px))
    bottom = img.crop((0, mid - overlap_px, w, h))

    tiles = []
    for tile, label_suffix in ((top, "top"), (bottom, "bottom")):
        buf = io.BytesIO()
        tile.save(buf, format="PNG")
        tiles.append({
            "page": page,
            "data": base64.b64encode(buf.getvalue()).decode("utf-8"),
            "mime": "image/png",
            "tile": label_suffix,
        })
    return tiles


def extract_pdf_pages(raw: bytes) -> list:
    """Extract each PDF page as {page_num, text, images}."""
    try:
        import fitz
    except ImportError:
        raise HTTPException(status_code=500, detail="PyMuPDF not installed. Run: pip install pymupdf")

    doc = fitz.open(stream=raw, filetype="pdf")
    pages = []

    for page_num, page in enumerate(doc):
        text = page.get_text("text").strip()

        # Render full page at 3x for better text legibility in dense documents
        mat = fitz.Matrix(3.0, 3.0)
        pix = page.get_pixmap(matrix=mat)
        img_bytes = pix.tobytes("png")
        b64 = base64.b64encode(img_bytes).decode("utf-8")
        # Tile the render into top/bottom halves so the LLM inspects each region
        images = _tile_image_b64(b64, page_num + 1)

        # Also extract each embedded image at native resolution so the LLM can
        # inspect each one individually at full quality (avoids missing small images
        # that are hard to read in the whole-page render).
        for img_info in page.get_images(full=True):
            xref = img_info[0]
            try:
                base_img = doc.extract_image(xref)
                img_data = base_img["image"]
                if len(img_data) > 4 * 1024 * 1024:  # skip images > 4 MB
                    continue
                ext = base_img.get("ext", "png").lower()
                if ext == "jpg":
                    ext = "jpeg"
                if ext not in ("png", "jpeg", "gif", "webp"):
                    ext = "png"
                emb_b64 = base64.b64encode(img_data).decode("utf-8")
                images.append({
                    "page": page_num + 1,
                    "data": emb_b64,
                    "mime": f"image/{ext}",
                    "embedded": True,
                })
            except Exception:
                pass

        # Collect font-size spans so the heading detector can decide whether
        # the title_case pass is worth running on this page.
        spans = []
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    txt = span.get("text", "").strip()
                    size = span.get("size", 0)
                    if txt and size:
                        spans.append({"text": txt, "size": size})

        pages.append({
            "page_num": page_num + 1,
            "text": text,
            "images": images,
            "spans": spans,
        })

    doc.close()
    return pages


def extract_pdf_zones(raw: bytes) -> tuple[list, list]:
    """
    Detect text blocks and image regions per page, draw colored boxes directly
    on each page render, and return (zones_per_page, annotated_preview_b64_list).

    Text blocks  → blue  boxes
    Image regions → orange boxes
    """
    try:
        import fitz
    except ImportError:
        raise HTTPException(status_code=500, detail="PyMuPDF not installed.")

    _TEXT_COLOR  = (52/255, 152/255, 219/255)   # blue
    _IMAGE_COLOR = (1.0,    140/255, 0.0)        # orange

    doc = fitz.open(stream=raw, filetype="pdf")
    zones_result = []
    previews = []
    page_count = len(doc)

    for page_num in range(page_count):
        page = doc[page_num]
        zones = []
        zone_id = 0

        # ── Text blocks ──────────────────────────────────────────────────────
        for block in page.get_text("blocks"):
            x0, y0, x1, y1, text, _block_no, block_type = block
            if block_type != 0 or not text.strip():
                continue
            rect = fitz.Rect(x0, y0, x1, y1)
            page.draw_rect(
                rect,
                color=_TEXT_COLOR,
                width=1.0,
                fill=_TEXT_COLOR,
                fill_opacity=0.10,
                overlay=True,
            )
            zones.append({
                "id": f"p{page_num+1}_text_{zone_id}",
                "type": "text",
                "label": "Text Block",
                "bbox": [round(x0), round(y0), round(x1), round(y1)],
            })
            zone_id += 1

        # ── Image regions ─────────────────────────────────────────────────────
        for img in page.get_images(full=True):
            xref = img[0]
            try:
                rects = page.get_image_rects(xref)
            except Exception:
                continue
            for rect in rects:
                page.draw_rect(
                    rect,
                    color=_IMAGE_COLOR,
                    width=1.5,
                    fill=_IMAGE_COLOR,
                    fill_opacity=0.10,
                    overlay=True,
                )
                zones.append({
                    "id": f"p{page_num+1}_img_{zone_id}",
                    "type": "image",
                    "label": "Image",
                    "bbox": [round(rect.x0), round(rect.y0), round(rect.x1), round(rect.y1)],
                })
                zone_id += 1

        # Fallback: if a page has no text AND no detectable images (e.g. the whole
        # page is a flattened/form-object image), mark the full page as an image zone.
        if not zones and not page.get_text().strip():
            pw, ph = page.rect.width, page.rect.height
            full_rect = fitz.Rect(0, 0, pw, ph)
            page.draw_rect(
                full_rect,
                color=_IMAGE_COLOR,
                width=1.5,
                fill=_IMAGE_COLOR,
                fill_opacity=0.08,
                overlay=True,
            )
            zones.append({
                "id": f"p{page_num+1}_img_{zone_id}",
                "type": "image",
                "label": "Image (full page)",
                "bbox": [0, 0, round(pw), round(ph)],
            })

        # ── Render annotated page to PNG ──────────────────────────────────────
        mat = fitz.Matrix(1.5, 1.5)
        pix = page.get_pixmap(matrix=mat)
        previews.append(base64.b64encode(pix.tobytes("png")).decode("utf-8"))
        zones_result.append({"page": page_num + 1, "zones": zones})

    doc.close()
    return zones_result, previews


def _strip_font_subset(name: str) -> str:
    """Strip subset prefix like 'ABCDEF+FontName' → 'FontName'."""
    return name.split("+", 1)[1] if "+" in name else name


def extract_pdf_fonts(raw: bytes) -> dict:
    """Return {font_name: {pages, page_count}} for all text fonts in the PDF."""
    try:
        import fitz
    except ImportError:
        raise HTTPException(status_code=500, detail="PyMuPDF not installed. Run: pip install pymupdf")

    doc = fitz.open(stream=raw, filetype="pdf")
    font_pages: dict = {}

    for page_num, page in enumerate(doc):
        for font in page.get_fonts(full=True):
            name = _strip_font_subset(font[3] or font[4] or "Unknown")
            font_pages.setdefault(name, set()).add(page_num + 1)

    doc.close()
    return {
        name: {"pages": sorted(pages), "page_count": len(pages)}
        for name, pages in sorted(font_pages.items())
    }


def extract_font_rects(raw: bytes, font_name: str, page_num: int) -> dict:
    """Return bounding boxes of all text spans using font_name on page_num (1-based).
    Only text spans are returned — image blocks are skipped by design."""
    try:
        import fitz
    except ImportError:
        raise HTTPException(status_code=500, detail="PyMuPDF not installed. Run: pip install pymupdf")

    doc = fitz.open(stream=raw, filetype="pdf")
    if page_num < 1 or page_num > len(doc):
        doc.close()
        return {"rects": [], "page_width": 612, "page_height": 792}

    page = doc[page_num - 1]
    page_width = page.rect.width
    page_height = page.rect.height

    rects = []
    blocks = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]
    for block in blocks:
        if block.get("type") != 0:   # 0 = text, 1 = image — skip images
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                if _strip_font_subset(span.get("font", "")) == font_name:
                    b = span["bbox"]
                    rects.append({"x0": b[0], "y0": b[1], "x1": b[2], "y1": b[3]})

    doc.close()
    return {"rects": rects, "page_width": page_width, "page_height": page_height}


def _extract_docx(raw: bytes) -> Tuple[str, str, list]:
    try:
        from docx import Document
    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx not installed. Run: pip install python-docx")

    doc = Document(io.BytesIO(raw))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    image_blocks = []

    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            try:
                img_bytes = rel.target_part.blob
                mime = rel.target_part.content_type
                b64 = base64.b64encode(img_bytes).decode("utf-8")
                image_blocks.append({"page": None, "data": b64, "mime": mime})
            except Exception:
                pass

    return "\n".join(paragraphs), "docx", image_blocks


def _extract_image(raw: bytes, filename: str) -> Tuple[str, str, list]:
    mime_map = {
        ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".jfif": "image/jpeg",
        ".png": "image/png", ".webp": "image/webp", ".gif": "image/gif",
    }
    ext = "." + filename.rsplit(".", 1)[-1]
    b64 = base64.b64encode(raw).decode("utf-8")
    # Tile tall images so the LLM focuses on each region
    tiles = _tile_image_b64(b64, page=1)
    return "", "image", tiles
def convert_to_pdf(raw: bytes, filename: str) -> bytes:
    import fitz
    text = ""
    if filename.lower().endswith(".docx"):
        from docx import Document
        import io
        try:
            doc = Document(io.BytesIO(raw))
            text = '\n'.join([p.text for p in doc.paragraphs if p.text.strip()])
        except Exception:
            text = "Failed to extract DOCX text."
    else:
        text = raw.decode("utf-8", errors="ignore")
        
    pdf_doc = fitz.open()
    lines = text.split('\n')
    lines_per_page = 45
    
    if not lines or not any(lines):
        # Empty doc
        page = pdf_doc.new_page(width=595, height=842)
        page.insert_textbox(fitz.Rect(50, 50, 545, 792), "[Empty Document]", fontsize=11, fontname="helv")
    else:
        for i in range(0, max(1, len(lines)), lines_per_page):
            page_lines = lines[i:i+lines_per_page]
            page_text = '\n'.join(page_lines)
            
            page = pdf_doc.new_page(width=595, height=842)
            rect = fitz.Rect(50, 50, 545, 792)
            page.insert_textbox(rect, page_text, fontsize=11, fontname="helv")
            
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()
    return pdf_bytes
